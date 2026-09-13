"""直接用 openpyxl 生成台账；只统计符合项，待确认单独列出。"""
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ["发布日期", "标题", "来源", "主题", "判断", "判断理由", "原文证据",
           "原文链接", "文章ID", "正文状态", "日期说明", "命中关键词"]


def add_sheet(workbook, name, headers, rows, widths):
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        values = []
        for value in row:
            if isinstance(value, str):
                if len(value) > 32767:
                    raise ValueError(f"{name} 第 {sheet.max_row + 1} 行超出 Excel 单元格文本上限；完整内容保留在 JSON 中")
                value = ILLEGAL_CHARACTERS_RE.sub("", value)
            values.append(value)
        sheet.append(values)
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # 外部文本一律按字面量保存，避免 Excel 公式注入。
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="315C4C")
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    return sheet


def article_rows(articles):
    for a in articles:
        hits = a.get("keyword_hits", {})
        words = list(dict.fromkeys(hits.get("title", []) + hits.get("content", [])))
        yield [a["date"], a["title"], "、".join(a["sources"]), a["category"], a["decision"],
               a["reason"], a["evidence"], a["url"], a["id"], a.get("content_status", ""),
               a.get("date_note", ""), "、".join(words)]


def export_excel(data, target):
    if data["semantic_state"] != "已完成" or data["summary"]["pending_count"]:
        raise ValueError("语义判断未完成，不导出成品")
    target = Path(target)
    if target.exists():
        raise ValueError(f"文件已存在，拒绝覆盖：{target}")
    wb = Workbook()
    wb.remove(wb.active)
    s = data["summary"]
    metrics = [
        ("流程", "区间内去重文章", s["candidate_count"]),
        ("流程", "关键词粗筛通过", s["coarse_count"]),
        ("流程", "关键词排除", s["keyword_excluded_count"]),
        ("流程", "已判断", s["judged_count"]),
        ("结论", "符合（正式统计）", s["accepted_count"]),
        ("结论", "不符合", s["rejected_count"]),
        ("结论", "待确认（不计入符合）", s["review_count"]),
        ("结论", "未判断", s["pending_count"]),
        ("覆盖", data["coverage_state"], s["issue_count"]),
    ]
    for label, key in [("按月（符合）", "by_month"), ("按主题（符合）", "by_category"), ("按来源（符合）", "by_source")]:
        metrics.extend((label, name, count) for name, count in s[key].items())
    add_sheet(wb, "统计", ["维度", "项目", "篇数"], metrics, [24, 60, 16])
    for name, articles in [("新闻台账", data["articles"]), ("待确认", data["review_articles"]), ("全部判断", data["evaluations"])]:
        sheet = add_sheet(wb, name, HEADERS, article_rows(articles), [14, 60, 28, 20, 14, 55, 70, 50, 25, 30, 30, 35])
        for row, a in enumerate(articles, 2):
            cell = sheet.cell(row, 8)
            cell.hyperlink = a["url"]
            cell.style = "Hyperlink"
    rows = []
    for source in data["sources"]:
        rows.append([source["source"], source["state"], source["pages"], source["checked"],
                     source["matched"], source["stop"], source["url"]])
        rows.extend([source["source"], "缺项", None, None, None, e["error"], e["url"]] for e in source["errors"])
    sheet = add_sheet(wb, "来源与缺项", ["来源", "状态", "扫描页数", "检查条目", "符合篇数", "说明", "链接"],
                      rows, [28, 26, 14, 14, 14, 70, 55])
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 7).hyperlink = sheet.cell(row, 7).value
        sheet.cell(row, 7).style = "Hyperlink"
    add_sheet(wb, "采集配置", ["配置项", "值"], [
        ["开始日期", data["start"]], ["结束日期", data["end"]], ["语义要求", data["topic"]],
        ["标题关键词", "、".join(data["keywords"]["title"])],
        ["全文关键词", "、".join(data["keywords"]["content"])],
        ["运行ID", data["run_id"]], ["生成时间", data["generated_at"]],
        ["统计口径", "正式统计只包含符合项；按文章 ID 去重；同篇文章可属于多个来源，来源数不可直接相加。"],
        ["覆盖边界", "仅限所选栏目与本次采集范围；来源缺项另列，不代表全校所有新闻。"],
    ], [24, 110])
    buffer = BytesIO()
    wb.save(buffer)
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("xb") as stream:
        stream.write(buffer.getvalue())
    return target
